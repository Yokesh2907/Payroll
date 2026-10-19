"""Payroll Master Application - multi-employee payroll grid with Excel import/export and PDF output.

Speed note: the grid is a single ttk.Treeview with ONE floating cell editor (instead of ~3,000 separate
Entry widgets), so the program opens, scrolls, saves and closes quickly even with hundreds of employees.
"""
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from tkinter import font as tkfont
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from datetime import datetime
import os, sys, subprocess, re, shutil, json

from payroll_calc import (OT_METHODS, LEGACY_OT_METHODS, DEFAULT_SETTINGS, D, WHOLE_FIELDS,
                          num, num_or_zero, compute, fmt)

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    OPENPYXL = True
except ImportError:
    OPENPYXL = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    REPORTLAB = True
except ImportError:
    REPORTLAB = False

APP_DIR = Path(__file__).resolve().parent
REFERENCE_XLSX = APP_DIR / 'Payroll_Reference.xlsx'
DEFAULT_XLSX = APP_DIR / 'Payroll_Data.xlsx'
BACKUP_DIR = APP_DIR / 'backups'
SESSION_FILE = APP_DIR / 'last_session.json'
MAX_BACKUPS = 10
SALARY_SHEETS = ['GOVINDARAJ SALARY', 'MD ARMAN SALARY', 'V BALAJI SALARY']
EMPLOYEE_MASTER_SHEET = 'DETAILS  FOR PAY ROLL'
ACTIVE_MEMBERS_PREFIX = 'activemembers'   # sheet is named ActiveMembers_<date>; the date changes monthly
ALL_TAB = '★ All employees (master + salary sheets)'

# ---------------------------------------------------------------- theme
BG, CARD, NAVY, NAVY2 = '#f2f5fa', '#ffffff', '#1b2a49', '#2a3d68'
ACCENT, GREEN, RED, AMBER = '#2563eb', '#16a34a', '#dc2626', '#d97706'
TEXT, MUTED, LINE = '#1f2937', '#64748b', '#d7deeb'
ROW_ODD, ROW_EVEN, ROW_BAD, ROW_WARN = '#ffffff', '#f3f7fd', '#fecaca', '#fef3c7'

# ---------------------------------------------------------------- columns
# Every data column (this is what is stored in Excel). (heading, key, grid width, type)
COLUMNS = [
    ('Employee Name','employee_name',180,'text'), ('Employee Code','employee_code',110,'text'),
    ('Company','company',120,'text'), ('Working Days','working_days',90,'num'), ('Paid Holidays','paid_holidays',90,'num'),
    ('Paid Days','paid_days',80,'calc'),
    ('Basic Salary','basic',105,'num'), ('Allowances','allowances',105,'num'), ('Other Earnings','other_earnings',105,'num'),
    ('OT Method','ot_method',250,'choice'), ('Multiplier','multiplier',85,'num'), ('Hours/Day','hours_per_day',85,'num'),
    ('OT Hours','ot_hours',80,'num'), ('OT Rate','ot_rate',90,'num'), ('OT Amount','ot_amount',100,'calc'),
    ('Gross Salary','gross',110,'calc'),
    ('Festival Advance','festival_advance',110,'num'), ('Festival Recovery','festival_recovery',115,'num'),
    ('Defer Festival Recovery?','defer_festival_recovery',90,'yesno'), ('Deferred Festival Recovery','deferred_festival_recovery',120,'calc'),
    ('Festival Previous Balance','festival_previous_balance',130,'num'), ('Festival Advance Balance','festival_balance',125,'calc'),
    ('General Advance Given','general_advance',120,'num'), ('General Advance Recovery','general_recovery',125,'num'),
    ('Defer Recovery?','defer_general_recovery',90,'yesno'), ('Deferred to Next Month','deferred_recovery',120,'calc'),
    ('Previous General Balance','previous_balance',130,'num'), ('General Advance Balance','general_balance',125,'calc'),
    ('Total Advance Deduction','advance_deduction',120,'calc'), ('Total Advance Balance','advance_balance',120,'calc'),
    ('PF (12% of Basic)','pf',105,'calc'), ('EPF part (3.67%)','epf',110,'calc'), ('EPS part (8.33%)','eps',110,'calc'),
    ('ESI Exempt','esi_exempt',80,'yesno'), ('ESI','esi',100,'calc'), ('Other Deductions','other_deductions',110,'num'),
    ('Total Deductions','total_deductions',115,'calc'), ('Net Salary','net_salary',110,'calc'),
]
KEYS = [c[1] for c in COLUMNS]
LABEL = {c[1]: c[0] for c in COLUMNS}
TYPE = {c[1]: c[3] for c in COLUMNS}
YESNO_KEYS = [k for k in KEYS if TYPE[k] == 'yesno']

# What the on-screen grid shows: a shorter, easier table. Advance details live in the Advance pop-up.
GRID = [
    ('Employee Name','employee_name',190), ('Code','employee_code',70),
    ('Work Days','working_days',80), ('Paid Hol.','paid_holidays',75), ('Paid Days','paid_days',75),
    ('Basic','basic',90), ('Allowances','allowances',90), ('Other Earnings','other_earnings',100),
    ('OT Method','ot_method',210), ('Mult.','multiplier',55), ('Hrs/Day','hours_per_day',65),
    ('OT Hrs','ot_hours',65), ('OT Rate','ot_rate',75), ('OT Amount','ot_amount',90), ('Gross','gross',100),
    ('Advance','advance_btn',100), ('Adv. Deduction','advance_deduction',105), ('Adv. Balance','advance_balance',100),
    ('PF','pf',85), ('EPF','epf',80), ('EPS','eps',80), ('ESI Exempt','esi_exempt',80), ('ESI','esi',80),
    ('Other Ded.','other_deductions',90), ('Total Ded.','total_deductions',100), ('Net Salary','net_salary',110),
]
GRID_KEYS = [g[1] for g in GRID]
TREE_COLS = ['idx'] + GRID_KEYS
GRID_TYPE = {k: ('button' if k == 'advance_btn' else TYPE[k]) for k in GRID_KEYS}
EDIT_KEYS = [k for k in GRID_KEYS if GRID_TYPE[k] in ('text', 'num', 'choice', 'yesno')]   # Enter/Tab walk through these
WHOLE_INPUTS = ('working_days', 'paid_holidays')    # entered as whole numbers

DEFAULTS = {k: ('26' if k == 'working_days' else '0') for k in KEYS}
DEFAULTS.update({'employee_name':'','employee_code':'','company':'','ot_method':OT_METHODS[0],
                 'multiplier':'1','hours_per_day':'8', **{k: 'No' for k in YESNO_KEYS}})

HEADER_ALIASES = {
    'employee name':'employee_name','employee':'employee_name','name':'employee_name',
    'employee code':'employee_code','code':'employee_code','emp code':'employee_code',
    'company':'company','company name':'company','working days':'working_days','w.days':'working_days','w days':'working_days',
    'paid holidays':'paid_holidays','ph':'paid_holidays','holiday':'paid_holidays','paid days':'paid_days',
    'basic salary':'basic','basic':'basic','allowances':'allowances','allowance':'allowances',
    'other earnings':'other_earnings','other earning':'other_earnings','ot hours':'ot_hours','o.t hrs':'ot_hours','ot hrs':'ot_hours',
    'ot method':'ot_method','multiplier':'multiplier','hours/day':'hours_per_day','hours per day':'hours_per_day','ot rate':'ot_rate',
    'ot amount':'ot_amount','ot':'ot_amount','gross salary':'gross','gross':'gross',
    'festival advance':'festival_advance','festival recovery':'festival_recovery','festival advance recovery':'festival_recovery',
    'defer festival recovery':'defer_festival_recovery','defer festival recovery?':'defer_festival_recovery',
    'deferred festival recovery':'deferred_festival_recovery','festival previous balance':'festival_previous_balance',
    'festival advance balance':'festival_balance','festival balance':'festival_balance',
    'general advance given':'general_advance','general advance':'general_advance','general recovery':'general_recovery',
    'general advance recovery':'general_recovery','defer recovery':'defer_general_recovery','defer recovery?':'defer_general_recovery',
    'deferred to next month':'deferred_recovery','previous general balance':'previous_balance','previous balance':'previous_balance',
    'general advance balance':'general_balance','general balance':'general_balance',
    'total advance deduction':'advance_deduction','total advance balance':'advance_balance',
    'other deductions':'other_deductions',
    'pf':'pf','pf (12% of basic)':'pf','epf':'epf','epf (3.67% of basic)':'epf','epf part (3.67%)':'epf',
    'eps':'eps','eps (8.33% of basic)':'eps','eps part (8.33%)':'eps',
    'esi':'esi','esi (1.5% of gross)':'esi','esi exempt':'esi_exempt',
    'total deductions':'total_deductions','net salary':'net_salary','net':'net_salary'
}
MONEY_NONZERO_KEYS = ['basic','allowances','other_earnings','ot_hours','ot_rate','festival_advance','festival_recovery',
                      'festival_previous_balance','general_advance','general_recovery','previous_balance','other_deductions']


def money(x, sym='₹'):
    d = num_or_zero(x)
    return f'{"-" if d < 0 else ""}{sym}{abs(d):,.2f}'

def plain(x):
    s = f'{num_or_zero(x):f}'
    return s.rstrip('0').rstrip('.') if '.' in s else s

def clean_key(x):
    return re.sub(r'\s+',' ',str(x or '').strip().lower().replace('_',' '))

def clean_person(x):
    value = clean_key(x)
    value = re.sub(r'\bmr\.?\s*','',value)
    return re.sub(r'\s+',' ',re.sub(r'[^a-z0-9]+',' ',value)).strip()

def norm_yesno(x):
    return 'Yes' if str(x).strip().lower() in ('yes','y','true','1') else 'No'

def to_whole(text):
    """'25.0' -> '25', '24.5' -> '25'. Non-numbers are returned unchanged."""
    try:
        return str(num(text).quantize(D('1'), rounding=ROUND_HALF_UP))
    except ValueError:
        return text

def make_record(values=None):
    """A complete employee record (all columns) from a partial dict, with tidy number formats."""
    values = values or {}
    rec = {}
    for k in KEYS:
        v = values.get(k, DEFAULTS[k])
        v = DEFAULTS[k] if v is None else str(v)
        if k == 'ot_method':
            v = LEGACY_OT_METHODS.get(v, v) if (v in LEGACY_OT_METHODS or v in OT_METHODS) else OT_METHODS[0]
        elif TYPE[k] == 'yesno':
            v = norm_yesno(v)
        elif TYPE[k] in ('num', 'calc'):
            try: v = plain(num(v)) if v.strip() != '' else '0'     # 12012.0 -> 12012 (text that is not a number is kept as typed)
            except ValueError: pass
        elif k == 'employee_code':
            v = re.sub(r'^(\d+)\.0+$', r'\1', v.strip())
        else:
            v = v.strip()
        rec[k] = v
    for k in WHOLE_INPUTS: rec[k] = to_whole(rec[k])
    return rec


def pdf_fonts():
    """Return (regular, bold, has_rupee_glyph). Uses a TTF that really contains the rupee sign."""
    if not REPORTLAB: return 'Helvetica','Helvetica-Bold',False
    cands = [(APP_DIR/'fonts'/'DejaVuSans.ttf', APP_DIR/'fonts'/'DejaVuSans-Bold.ttf'),
             ('C:/Windows/Fonts/arial.ttf','C:/Windows/Fonts/arialbd.ttf'),
             ('C:/Windows/Fonts/segoeui.ttf','C:/Windows/Fonts/segoeuib.ttf'),
             ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'),
             ('/System/Library/Fonts/Supplemental/Arial.ttf','/System/Library/Fonts/Supplemental/Arial Bold.ttf')]
    for reg, bold in cands:
        try:
            if not (Path(reg).exists() and Path(bold).exists()): continue
            t = TTFont('PayFont', str(reg))
            if 0x20B9 not in t.face.charToGlyph: continue
            pdfmetrics.registerFont(t); pdfmetrics.registerFont(TTFont('PayFont-Bold', str(bold)))
            return 'PayFont','PayFont-Bold',True
        except Exception:
            continue
    return 'Helvetica','Helvetica-Bold',False


def pick_font_family():
    try:
        fams = set(tkfont.families())
    except Exception:
        return 'TkDefaultFont'
    for f in ('Segoe UI', 'Helvetica Neue', 'Noto Sans', 'DejaVu Sans', 'Arial'):
        if f in fams: return f
    return tkfont.nametofont('TkDefaultFont').actual('family')


# =====================================================================================
#  Pop-up: choose Excel file + tab
# =====================================================================================
class OpenDialog(tk.Toplevel):
    """Asks which Excel file and which tab (sheet) to work on. result = (path, sheet) | ('blank', None) | None."""
    HINTS = [('payroll data', 'latest saved payroll'), ('payroll ', 'saved monthly payroll'),
             (clean_key(EMPLOYEE_MASTER_SHEET), 'employee master list'), (ACTIVE_MEMBERS_PREFIX, 'PF active members list')]

    def __init__(self, app, path, sheet):
        super().__init__(app)
        self.app = app; self.result = None; self.names = []
        self.title('Open Excel - choose file and tab'); self.configure(bg=BG); self.transient(app)
        self.geometry('680x520'); self.minsize(560, 440)
        head = tk.Frame(self, bg=NAVY); head.pack(fill='x')
        tk.Label(head, text='Which Excel do you want to open?', bg=NAVY, fg='white', font=(app.fam, 15, 'bold')).pack(anchor='w', padx=18, pady=(14, 2))
        tk.Label(head, text='Step 1: pick the file    Step 2: pick the tab (sheet)', bg=NAVY, fg='#c7d2ea', font=(app.fam, 10)).pack(anchor='w', padx=18, pady=(0, 14))
        body = ttk.Frame(self, padding=(18, 14)); body.pack(fill='both', expand=True)
        ttk.Label(body, text='1. Excel file', style='Bold.TLabel').pack(anchor='w')
        row = ttk.Frame(body); row.pack(fill='x', pady=(4, 12))
        self.path = tk.StringVar(value=str(path or ''))
        e = ttk.Entry(row, textvariable=self.path); e.pack(side='left', fill='x', expand=True)
        e.bind('<Return>', lambda ev: self.load_sheets())
        ttk.Button(row, text='Browse...', command=self.browse).pack(side='left', padx=(8, 0))
        ttk.Label(body, text='2. Tab (sheet) to work on', style='Bold.TLabel').pack(anchor='w')
        lf = ttk.Frame(body); lf.pack(fill='both', expand=True, pady=(4, 6))
        self.lb = tk.Listbox(lf, exportselection=False, activestyle='none', font=(app.fam, 11), relief='solid', borderwidth=1,
                             highlightthickness=0, selectbackground=ACCENT, selectforeground='white')
        sb = ttk.Scrollbar(lf, orient='vertical', command=self.lb.yview); self.lb.configure(yscrollcommand=sb.set)
        self.lb.pack(side='left', fill='both', expand=True); sb.pack(side='left', fill='y')
        self.lb.bind('<Double-Button-1>', lambda ev: self.ok()); self.lb.bind('<Return>', lambda ev: self.ok())
        self.note = tk.StringVar(value='')
        ttk.Label(body, textvariable=self.note, style='Muted.TLabel', wraplength=620).pack(anchor='w')
        btns = ttk.Frame(self, padding=(18, 4, 18, 16)); btns.pack(fill='x')
        ttk.Button(btns, text='Open', style='Primary.TButton', command=self.ok).pack(side='right')
        ttk.Button(btns, text='Start with blank sheet', command=self.blank).pack(side='right', padx=8)
        self.want_sheet = sheet
        self.protocol('WM_DELETE_WINDOW', self.destroy)
        self.bind('<Escape>', lambda ev: self.destroy())
        self.load_sheets()
        self.update_idletasks()
        x = app.winfo_rootx() + max(0, (app.winfo_width() - self.winfo_width()) // 2)
        y = app.winfo_rooty() + max(0, (app.winfo_height() - self.winfo_height()) // 3)
        self.geometry(f'+{x}+{y}')
        self.lb.focus_set()
        self.wait_visibility(); self.grab_set()

    def browse(self):
        p = filedialog.askopenfilename(parent=self, initialdir=str(Path(self.path.get()).parent) if self.path.get() else str(APP_DIR),
                                       filetypes=[('Excel files', '*.xlsx *.xlsm *.xltx *.xltm *.xls'), ('All files', '*.*')])
        if p: self.path.set(p); self.load_sheets()

    def hint(self, name):
        k = clean_key(name)
        for pre, text in self.HINTS:
            if k.startswith(pre): return text
        return 'single-employee salary slip' if name in SALARY_SHEETS else ''

    def load_sheets(self):
        self.lb.delete(0, 'end'); self.names = []
        p = self.path.get().strip().strip('"')
        if not p: self.note.set('Click Browse to choose an Excel file.'); return
        if not Path(p).exists(): self.note.set('That file was not found. Click Browse to choose another.'); return
        if p.lower().endswith('.xls'):
            self.note.set('Older .xls file: it will be converted (needs LibreOffice) when you press Open.')
            self.names = [None]; self.lb.insert('end', '(all tabs are read after conversion)'); self.lb.selection_set(0); return
        try:
            wb = load_workbook(p, read_only=True, data_only=True); sheets = list(wb.sheetnames); wb.close()
        except Exception as e:
            self.note.set(f'Could not read this file: {e}'); return
        self.names = [ALL_TAB] + sheets
        self.lb.insert('end', ALL_TAB)
        for n in sheets:
            h = self.hint(n); self.lb.insert('end', f'{n}      -  {h}' if h else n)
        pick = 0
        if self.want_sheet in self.names: pick = self.names.index(self.want_sheet)
        elif 'Payroll Data' in self.names: pick = self.names.index('Payroll Data')
        self.lb.selection_set(pick); self.lb.see(pick)
        self.note.set(f'{len(sheets)} tab(s) found. Double-click a tab to open it.')

    def ok(self):
        p = self.path.get().strip().strip('"')
        sel = self.lb.curselection()
        if not p or not Path(p).exists(): messagebox.showwarning('Choose a file', 'Please choose an Excel file first.', parent=self); return
        if not self.names or not sel: messagebox.showwarning('Choose a tab', 'Please select a tab (sheet).', parent=self); return
        self.result = (p, self.names[sel[0]]); self.destroy()

    def blank(self):
        self.result = ('blank', None); self.destroy()


# =====================================================================================
#  Pop-up: advance details of one employee
# =====================================================================================
class AdvanceDialog(tk.Toplevel):
    """General + festival advance for ONE employee: given, recovery, balance, defer / deferred to next month."""
    SECTIONS = [
        ('General Advance', 'previous_balance', 'general_advance', 'general_recovery', 'defer_general_recovery',
         'deferred_recovery', 'general_balance'),
        ('Festival Advance', 'festival_previous_balance', 'festival_advance', 'festival_recovery', 'defer_festival_recovery',
         'deferred_festival_recovery', 'festival_balance'),
    ]

    def __init__(self, app, iid):
        super().__init__(app)
        self.app, self.iid = app, iid
        d = app.data[iid]; who = d['employee_name'] or d['employee_code'] or 'employee'
        self.title(f'Advance - {who}'); self.configure(bg=BG); self.transient(app); self.resizable(False, False)
        head = tk.Frame(self, bg=NAVY); head.pack(fill='x')
        tk.Label(head, text='Advance details', bg=NAVY, fg='white', font=(app.fam, 15, 'bold')).pack(anchor='w', padx=18, pady=(12, 0))
        tk.Label(head, text=f'{who}   |   Code {d["employee_code"] or "-"}   |   {app.month.get()}', bg=NAVY, fg='#c7d2ea',
                 font=(app.fam, 10)).pack(anchor='w', padx=18, pady=(0, 12))
        body = ttk.Frame(self, padding=14); body.pack(fill='both', expand=True)
        self.vars, self.defer, self.out, self.order = {}, {}, {}, []
        for col, (title, prev, given, rec, defer, deferred, bal) in enumerate(self.SECTIONS):
            card = ttk.LabelFrame(body, text=f'  {title}  ', padding=12); card.grid(row=0, column=col, padx=7, sticky='nsew')
            body.columnconfigure(col, weight=1)
            for r, (label, key) in enumerate([('Previous balance (brought forward)', prev), ('Advance given this month', given),
                                              ('Recovery this month', rec)]):
                ttk.Label(card, text=label).grid(row=r, column=0, sticky='w', pady=4)
                v = tk.StringVar(value=d[key]); self.vars[key] = v; v.trace_add('write', self.refresh)
                e = ttk.Entry(card, textvariable=v, width=14, justify='right'); e.grid(row=r, column=1, padx=(14, 0), pady=4)
                self.order.append(e)
            dv = tk.BooleanVar(value=d[defer] == 'Yes'); self.defer[defer] = dv; dv.trace_add('write', self.refresh)
            ttk.Label(card, text='Defer recovery to next month?').grid(row=3, column=0, sticky='w', pady=4)
            cb = ttk.Checkbutton(card, text='Yes, do not deduct now', variable=dv); cb.grid(row=3, column=1, padx=(14, 0), sticky='w')
            self.order.append(cb)
            ttk.Separator(card).grid(row=4, column=0, columnspan=2, sticky='ew', pady=8)
            ttk.Label(card, text='Deferred to next month').grid(row=5, column=0, sticky='w', pady=3)
            ttk.Label(card, text='Balance after this month', style='Bold.TLabel').grid(row=6, column=0, sticky='w', pady=3)
            a = ttk.Label(card, text='', anchor='e', width=14); a.grid(row=5, column=1, padx=(14, 0), sticky='e')
            b = ttk.Label(card, text='', anchor='e', width=14, style='Bold.TLabel'); b.grid(row=6, column=1, padx=(14, 0), sticky='e')
            self.out[deferred], self.out[bal] = a, b
        foot = ttk.Frame(self, padding=(14, 0, 14, 4)); foot.pack(fill='x')
        self.summary = tk.StringVar(); ttk.Label(foot, textvariable=self.summary, style='Bold.TLabel').pack(anchor='w')
        self.msg = tk.StringVar(); self.msg_label = ttk.Label(foot, textvariable=self.msg, style='Muted.TLabel'); self.msg_label.pack(anchor='w', pady=(2, 0))
        btns = ttk.Frame(self, padding=(14, 6, 14, 14)); btns.pack(fill='x')
        self.save_btn = ttk.Button(btns, text='Save advance details', style='Primary.TButton', command=self.save); self.save_btn.pack(side='right')
        ttk.Button(btns, text='Cancel', command=self.destroy).pack(side='right', padx=8)
        ttk.Label(btns, text='Enter = next box   |   Esc = cancel', style='Muted.TLabel').pack(side='left')
        for i, w in enumerate(self.order):
            w.bind('<Return>', lambda ev, i=i: self.next(i))
        self.bind('<Escape>', lambda ev: self.destroy())
        self.refresh()
        self.update_idletasks()
        x = app.winfo_rootx() + max(0, (app.winfo_width() - self.winfo_width()) // 2)
        y = app.winfo_rooty() + max(0, (app.winfo_height() - self.winfo_height()) // 3)
        self.geometry(f'+{x}+{y}')
        self.order[0].focus_set(); self.order[0].selection_range(0, 'end')
        self.wait_visibility(); self.grab_set()

    def next(self, i):
        if i + 1 < len(self.order):
            w = self.order[i + 1]; w.focus_set()
            if isinstance(w, ttk.Entry): w.selection_range(0, 'end')
        else:
            self.save()
        return 'break'

    def current(self):
        tmp = dict(self.app.data[self.iid])
        for k, v in self.vars.items(): tmp[k] = v.get().strip() or '0'
        for k, v in self.defer.items(): tmp[k] = 'Yes' if v.get() else 'No'
        return tmp

    def refresh(self, *_):
        try:
            calc, warns = compute(self.current(), self.app.settings())
        except ValueError as e:
            self.msg.set(f'Please enter numbers only ({e})'); self.save_btn.state(['disabled'])
            self.msg_label.configure(style='Bad.TLabel'); return
        self.save_btn.state(['!disabled'])
        for k, lab in self.out.items(): lab.configure(text=money(calc[k]))
        self.summary.set(f'Deducted from salary this month: {money(calc["advance_deduction"])}      '
                         f'Total advance balance: {money(calc["advance_balance"])}')
        bad = [m for lvl, m, f in warns if f in ('general_balance', 'festival_balance')]
        self.msg.set(' '.join(bad) if bad else 'Tip: tick "Defer" to skip this month\'s deduction - the amount shows under "Deferred to next month".')
        self.msg_label.configure(style='Warn.TLabel' if bad else 'Muted.TLabel')

    def save(self):
        if 'disabled' in self.save_btn.state(): return
        d = self.app.data[self.iid]
        for k, v in self.current().items():
            if k in self.vars or k in self.defer: d[k] = plain(v) if k not in self.defer else v
        self.app.advance_saved(self.iid)
        self.destroy()


# =====================================================================================
#  Main application
# =====================================================================================
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title('Payroll Master')
        sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
        width, height = min(1600, max(900, sw - 50)), min(900, max(600, sh - 90))
        self.geometry(f'{width}x{height}+10+10'); self.minsize(min(850, sw), min(520, sh))
        self.fam = pick_font_family()
        self.data = {}              # iid -> full employee record (dict of strings)
        self.flags = {}             # iid -> 'bad' | 'warn' | None
        self.next_id = 0
        self.warnings = []; self.error_rows = []
        self.excel_path = None; self.sheet_name = ''
        self.dirty = False
        self._loading = False; self._recalc_job = None
        self.editor = None; self.cur = None
        self.build_style(); self.build_ui()
        self.protocol('WM_DELETE_WINDOW', self.on_exit)
        self.bind_all('<Control-s>', lambda e: self.run(self.save_to_excel))
        self.bind_all('<Control-o>', lambda e: self.run(self.ask_open))
        self.set_rows([])
        self.after(150, lambda: self.ask_open(startup=True))

    # ------------------------------------------------------------------ style
    def build_style(self):
        st = ttk.Style(self); st.theme_use('clam'); f = self.fam
        self.option_add('*Font', (f, 10))
        st.configure('.', font=(f, 10), background=BG, foreground=TEXT, bordercolor=LINE, lightcolor=BG, darkcolor=BG, troughcolor='#e3e8f2')
        st.configure('TFrame', background=BG); st.configure('Card.TFrame', background=CARD)
        st.configure('TLabel', background=BG, foreground=TEXT)
        st.configure('Bold.TLabel', font=(f, 10, 'bold')); st.configure('Muted.TLabel', foreground=MUTED)
        st.configure('Bad.TLabel', foreground=RED); st.configure('Warn.TLabel', foreground=AMBER, font=(f, 10, 'bold'))
        st.configure('TLabelframe', background=BG, bordercolor=LINE, relief='solid')
        st.configure('TLabelframe.Label', background=BG, foreground=NAVY, font=(f, 10, 'bold'))
        st.configure('TEntry', fieldbackground='white', padding=4, bordercolor=LINE)
        st.configure('TCombobox', fieldbackground='white', padding=4, bordercolor=LINE, arrowsize=14)
        st.map('TCombobox', fieldbackground=[('readonly', 'white')])
        st.configure('TCheckbutton', background=BG)
        st.map('TCheckbutton', background=[('active', BG)])
        st.configure('TButton', background='#e2e8f5', foreground=NAVY, padding=(12, 7), relief='flat', borderwidth=0, font=(f, 10, 'bold'))
        st.map('TButton', background=[('active', '#cfd9ef'), ('disabled', '#eceff5')], foreground=[('disabled', '#9aa5b8')])
        for name, bg, hov in [('Primary', ACCENT, '#1d4ed8'), ('Success', GREEN, '#15803d'), ('Danger', RED, '#b91c1c'), ('Dark', NAVY2, '#3b5088')]:
            st.configure(f'{name}.TButton', background=bg, foreground='white')
            st.map(f'{name}.TButton', background=[('active', hov), ('disabled', '#a5b4cf')], foreground=[('disabled', 'white')])
        st.configure('TMenubutton', background='#e2e8f5', foreground=NAVY, padding=(12, 7), relief='flat', borderwidth=0, font=(f, 10, 'bold'), arrowcolor=NAVY)
        st.map('TMenubutton', background=[('active', '#cfd9ef')])
        st.configure('Dark.TMenubutton', background=NAVY2, foreground='white', padding=(12, 7), relief='flat', borderwidth=0, font=(f, 10, 'bold'), arrowcolor='white')
        st.map('Dark.TMenubutton', background=[('active', '#3b5088')])
        st.configure('Payroll.Treeview', background=ROW_ODD, fieldbackground=ROW_ODD, foreground=TEXT, rowheight=30,
                     bordercolor=LINE, borderwidth=1, font=(f, 10))
        st.configure('Payroll.Treeview.Heading', background=NAVY2, foreground='white', font=(f, 10, 'bold'), relief='flat', padding=(4, 8))
        st.map('Payroll.Treeview.Heading', background=[('active', '#3b5088')])
        # Tk 8.6.9 bug workaround so row colours (tags) show up
        st.map('Payroll.Treeview', background=[('selected', ACCENT)], foreground=[('selected', 'white')])

    # ------------------------------------------------------------------ UI
    def month_choices(self):
        """A practical month/year dropdown spanning years around the current date."""
        from datetime import date
        now = date.today().replace(day=1); start = now.year - 3
        return [datetime(y, m, 1).strftime('%B %Y') for y in range(start, now.year + 4) for m in range(1, 13)]

    def build_ui(self):
        f = self.fam
        # ---- title bar with live totals
        head = tk.Frame(self, bg=NAVY); head.pack(fill='x')
        left = tk.Frame(head, bg=NAVY); left.pack(side='left', padx=18, pady=10)
        tk.Label(left, text='PAYROLL MASTER', bg=NAVY, fg='white', font=(f, 18, 'bold')).pack(anchor='w')
        self.file_label = tk.StringVar(value='No Excel opened yet')
        tk.Label(left, textvariable=self.file_label, bg=NAVY, fg='#c7d2ea', font=(f, 10)).pack(anchor='w')
        right = tk.Frame(head, bg=NAVY); right.pack(side='right', padx=14, pady=8)
        self.sum_vars = {}
        for key, cap, color in [('count', 'EMPLOYEES', '#c7d2ea'), ('gross', 'GROSS', '#93c5fd'), ('ded', 'DEDUCTIONS', '#fca5a5'), ('net', 'NET PAY', '#86efac')]:
            card = tk.Frame(right, bg=NAVY2); card.pack(side='left', padx=4)
            tk.Label(card, text=cap, bg=NAVY2, fg='#c7d2ea', font=(f, 8, 'bold')).pack(anchor='w', padx=12, pady=(6, 0))
            v = tk.StringVar(value='-'); self.sum_vars[key] = v
            tk.Label(card, textvariable=v, bg=NAVY2, fg=color, font=(f, 13, 'bold')).pack(anchor='w', padx=12, pady=(0, 6))

        root = ttk.Frame(self, padding=(12, 10, 12, 8)); root.pack(fill='both', expand=True)

        # ---- settings (month + working days always visible, the rest on demand)
        top = ttk.LabelFrame(root, text='  Payroll settings  ', padding=(10, 6)); top.pack(fill='x')
        self.month = tk.StringVar(value=datetime.now().strftime('%B %Y'))
        self.company_working_days = tk.StringVar(value='26')
        self.prorate = tk.BooleanVar(value=DEFAULT_SETTINGS['prorate'])
        self.esi_rate = tk.StringVar(value=str(DEFAULT_SETTINGS['esi_rate']))
        self.esi_ceiling = tk.StringVar(value=str(DEFAULT_SETTINGS['esi_ceiling']))
        self.pf_ceiling = tk.StringVar(value=str(DEFAULT_SETTINGS['pf_ceiling']))
        self.round_whole = tk.BooleanVar(value=DEFAULT_SETTINGS['round_whole'])
        self.esi_excl_allow = tk.BooleanVar(value=DEFAULT_SETTINGS['esi_excl_allow'])
        r1 = ttk.Frame(top); r1.pack(fill='x')
        ttk.Label(r1, text='Payroll month').pack(side='left')
        self.month_box = ttk.Combobox(r1, textvariable=self.month, width=16, state='readonly', values=self.month_choices())
        self.month_box.pack(side='left', padx=(6, 18)); self.month_box.bind('<<ComboboxSelected>>', self.on_month_selected)
        ttk.Label(r1, text='Company working days').pack(side='left')
        ttk.Entry(r1, textvariable=self.company_working_days, width=6, justify='center').pack(side='left', padx=(6, 18))
        ttk.Checkbutton(r1, text='Pay pro-rata by attendance (Working Days + Paid Holidays)', variable=self.prorate).pack(side='left')
        self.more_btn = ttk.Button(r1, text='More settings  ▾', command=self.toggle_more); self.more_btn.pack(side='right')
        self.more = ttk.Frame(top)
        ttk.Label(self.more, text='ESI rate %').grid(row=0, column=0, padx=(0, 4), pady=6, sticky='w')
        ttk.Entry(self.more, textvariable=self.esi_rate, width=7, justify='center').grid(row=0, column=1, padx=(0, 16))
        ttk.Label(self.more, text='ESI wage ceiling (0 = none)').grid(row=0, column=2, padx=(0, 4), sticky='w')
        ttk.Entry(self.more, textvariable=self.esi_ceiling, width=8, justify='center').grid(row=0, column=3, padx=(0, 16))
        ttk.Label(self.more, text='PF wage ceiling (0 = none)').grid(row=0, column=4, padx=(0, 4), sticky='w')
        ttk.Entry(self.more, textvariable=self.pf_ceiling, width=8, justify='center').grid(row=0, column=5, padx=(0, 16))
        ttk.Checkbutton(self.more, text='Round deductions to whole ₹', variable=self.round_whole).grid(row=1, column=0, columnspan=3, sticky='w')
        ttk.Checkbutton(self.more, text='Exclude allowances from ESI base', variable=self.esi_excl_allow).grid(row=1, column=3, columnspan=3, sticky='w')
        for v in (self.company_working_days, self.prorate, self.esi_rate, self.esi_ceiling, self.pf_ceiling, self.round_whole, self.esi_excl_allow):
            v.trace_add('write', self.schedule_recalc)

        # ---- toolbar
        toolbar = ttk.Frame(root); toolbar.pack(fill='x', pady=(10, 6))
        def btn(text, cmd, style='', side='left'):
            ttk.Button(toolbar, text=text, style=(style + '.TButton') if style else 'TButton', command=lambda: self.run(cmd)).pack(side=side, padx=(0, 6))
        def menu_btn(text, items, style, side='left'):
            m = tk.Menu(toolbar, tearoff=0, font=(f, 10), bg='white', activebackground=ACCENT, activeforeground='white', relief='flat', bd=1)
            for label, cmd, accel in items: m.add_command(label=label, accelerator=accel, command=lambda c=cmd: self.run(c))
            ttk.Menubutton(toolbar, text=text, style=style, menu=m).pack(side=side, padx=(0, 6))
        menu_btn('Excel', [('Open Excel / Tab...', self.ask_open, 'Ctrl+O'), ('Save to Excel', self.save_to_excel, 'Ctrl+S')], 'Dark.TMenubutton')
        ttk.Separator(toolbar, orient='vertical').pack(side='left', fill='y', padx=8)
        btn('Add Employee', self.add_employee, 'Primary'); btn('Advance', self.open_selected_advance, 'Primary')
        btn('Calculate All', self.calculate_all); btn('Delete Selected', self.delete_selected, 'Danger'); btn('Clear All', self.clear_all, 'Danger')
        btn('Exit', self.on_exit, 'Dark', 'right')
        menu_btn('PDF / Print', [('Generate PDF...', self.pdf_all, ''), ('Print...', self.print_all, '')], 'TMenubutton', 'right')

        self.banner = tk.StringVar(value='Click any white cell to type. Enter = next cell.  Click [ Advance ] on a row for advance / festival advance details.')
        ttk.Label(root, textvariable=self.banner, style='Muted.TLabel').pack(fill='x', pady=(0, 4))

        # ---- grid
        frame = ttk.Frame(root); frame.pack(fill='both', expand=True)
        self.tree = ttk.Treeview(frame, columns=TREE_COLS, show='headings', selectmode='browse', style='Payroll.Treeview')
        self.vs = ttk.Scrollbar(frame, orient='vertical', command=self.tree.yview)
        self.hs = ttk.Scrollbar(frame, orient='horizontal', command=self.tree.xview)
        self.tree.configure(yscrollcommand=self._yscroll, xscrollcommand=self._xscroll)
        self.tree.grid(row=0, column=0, sticky='nsew'); self.vs.grid(row=0, column=1, sticky='ns'); self.hs.grid(row=1, column=0, sticky='ew')
        frame.rowconfigure(0, weight=1); frame.columnconfigure(0, weight=1)
        self.tree.heading('idx', text='#'); self.tree.column('idx', width=44, minwidth=30, stretch=False, anchor='center')
        for label, key, w in GRID:
            w = max(w, len(label) * 9 + 30)       # never clip a heading
            anchor = 'w' if GRID_TYPE[key] in ('text', 'choice') else ('center' if GRID_TYPE[key] in ('yesno', 'button') else 'e')
            self.tree.heading(key, text=label, anchor='center'); self.tree.column(key, width=w, minwidth=40, stretch=False, anchor=anchor)
        self.tree.tag_configure('odd', background=ROW_ODD); self.tree.tag_configure('even', background=ROW_EVEN)
        self.tree.tag_configure('bad', background=ROW_BAD); self.tree.tag_configure('warn', background=ROW_WARN)
        self.tree.bind('<Button-1>', self.on_click)
        self.tree.bind('<Return>', self.on_tree_return)
        self.tree.bind('<<TreeviewSelect>>', self.on_select)
        self.tree.bind('<ButtonRelease-1>', lambda e: self.after_idle(self._reposition_editor))

        self.status = tk.StringVar(value='Ready.')
        bar = tk.Frame(self, bg=NAVY2); bar.pack(fill='x', side='bottom')
        tk.Label(bar, textvariable=self.status, bg=NAVY2, fg='white', anchor='w', font=(f, 10)).pack(fill='x', padx=12, pady=5)

    def toggle_more(self):
        if self.more.winfo_ismapped():
            self.more.pack_forget(); self.more_btn.configure(text='More settings  ▾')
        else:
            self.more.pack(fill='x'); self.more_btn.configure(text='Hide settings  ▴')

    def run(self, fn):
        """Run a toolbar action after saving whatever cell is being typed."""
        self.commit_editor()
        fn()

    def on_month_selected(self, event=None):
        self.dirty = True
        self.status.set(f"Payroll period: {self.month.get()}. Save to Excel to keep this month as its own tab.")

    def set_title(self):
        self.title(('* ' if self.dirty else '') + 'Payroll Master' + (f' - {self.excel_path.name}' if self.excel_path else ''))

    def mark_dirty(self):
        if not self.dirty and not self._loading:
            self.dirty = True; self.set_title()

    # ------------------------------------------------------------------ settings
    def settings(self):
        def dec(var, label):
            try: return num(var.get())
            except ValueError: raise ValueError(f'{label}: please enter a number.')
        return {'company_days': dec(self.company_working_days,'Company Working Days'),
                'prorate': self.prorate.get(), 'esi_rate': dec(self.esi_rate,'ESI rate'),
                'esi_ceiling': dec(self.esi_ceiling,'ESI wage ceiling'), 'pf_ceiling': dec(self.pf_ceiling,'PF wage ceiling'),
                'round_whole': self.round_whole.get(), 'esi_excl_allow': self.esi_excl_allow.get()}

    def settings_rows(self):
        return [('Payroll Month', self.month.get()), ('Company Working Days', self.company_working_days.get()),
                ('Prorate By Attendance', 'Yes' if self.prorate.get() else 'No'), ('ESI Rate %', self.esi_rate.get()),
                ('ESI Wage Ceiling', self.esi_ceiling.get()), ('PF Wage Ceiling', self.pf_ceiling.get()),
                ('Round To Whole Rupees', 'Yes' if self.round_whole.get() else 'No'),
                ('Exclude Allowances From ESI Base', 'Yes' if self.esi_excl_allow.get() else 'No'),
                ('OT Method / Multiplier / Hours per Day', 'Configured separately in each employee row')]

    def apply_settings(self, s):
        yes = lambda v: str(v).strip().lower() in ('yes','true','1','y')
        for key, var in [('payroll month', self.month), ('company working days', self.company_working_days),
                         ('esi rate %', self.esi_rate), ('esi wage ceiling', self.esi_ceiling), ('pf wage ceiling', self.pf_ceiling)]:
            if key in s and s[key] not in (None, ''): var.set(str(s[key]))
        self.ensure_month_listed()
        for key, var in [('prorate by attendance', self.prorate), ('round to whole rupees', self.round_whole),
                         ('exclude allowances from esi base', self.esi_excl_allow)]:
            if key in s and s[key] not in (None, ''): var.set(yes(s[key]))

    def ensure_month_listed(self):
        if self.month.get() not in self.month_box['values']:
            self.month_box.configure(values=tuple(self.month_box['values']) + (self.month.get(),))

    def schedule_recalc(self, *_):
        if self._loading: return
        self.mark_dirty()
        if self._recalc_job: self.after_cancel(self._recalc_job)
        self._recalc_job = self.after(400, lambda: self.calculate_all(False))

    # ------------------------------------------------------------------ opening an Excel file / tab
    def remembered(self):
        try: return json.loads(SESSION_FILE.read_text(encoding='utf-8'))
        except Exception: return {}

    def remember(self, path, sheet):
        try: SESSION_FILE.write_text(json.dumps({'path': str(path), 'sheet': sheet}), encoding='utf-8')
        except Exception: pass

    def ask_open(self, startup=False):
        if not OPENPYXL:
            messagebox.showerror('Missing Package', 'Install openpyxl:  pip install openpyxl'); return
        if not startup and self.dirty and self.active_ids():
            if not messagebox.askyesno('Unsaved changes', 'Opening another file/tab replaces the employees on screen.\nUnsaved changes will be lost. Continue?'):
                return
        last = self.remembered()
        path = last.get('path') if last.get('path') and Path(last['path']).exists() else None
        if not path:
            if not DEFAULT_XLSX.exists() and REFERENCE_XLSX.exists():
                try: shutil.copy2(REFERENCE_XLSX, DEFAULT_XLSX)
                except OSError: pass
            path = DEFAULT_XLSX if DEFAULT_XLSX.exists() else ''
        dlg = OpenDialog(self, path, last.get('sheet'))
        self.wait_window(dlg)
        res = dlg.result
        if res is None:
            if startup: self.status.set('No file opened. Click "Open Excel / Tab" to choose one, or just start typing.')
            return
        if res[0] == 'blank':
            self.set_rows([]); self.excel_path = None; self.sheet_name = ''; self.dirty = False; self.set_title()
            self.file_label.set('New blank sheet (Save asks where to store it)'); self.status.set('Blank sheet.'); self.update_summary(); return
        self.load_from(*res)

    def load_from(self, path, sheet):
        self.config(cursor='watch'); self.update_idletasks()
        try:
            path = Path(path); is_xls = path.suffix.lower() == '.xls'
            if is_xls:
                converted = path.with_suffix('.payroll_import.xlsx')
                if not self.convert_legacy_xls(path, converted):
                    messagebox.showerror('Legacy Excel', 'Could not read .xls. Please open it in Excel/LibreOffice and save it as .xlsx, then open again.'); return
                path = converted
            wb = load_workbook(path, data_only=True, read_only=True)
            try:
                month = None
                if sheet is None or sheet == ALL_TAB:
                    src, own = wb, False
                    if path.resolve() == DEFAULT_XLSX.resolve() and REFERENCE_XLSX.exists():
                        # the original workbook keeps the cached formula results of the salary sheets
                        src, own = load_workbook(REFERENCE_XLSX, data_only=True, read_only=True), True
                    data = self.read_all_reference_employees(src)
                    if own: src.close()
                    sheet = ALL_TAB
                elif sheet == EMPLOYEE_MASTER_SHEET:
                    recs, _ = self.read_master(wb); data = list(recs.values())
                elif clean_key(sheet).startswith(ACTIVE_MEMBERS_PREFIX):
                    recs, master_cos = self.read_master(wb) if EMPLOYEE_MASTER_SHEET in wb.sheetnames else ({}, {})
                    data = list(self.read_active(wb, {}, master_cos).values())
                else:
                    ws = wb[sheet]
                    if clean_key(sheet).startswith('payroll') or self.looks_like_export(ws):
                        data = self.read_tabular_sheet(ws); month = self.sheet_month(ws)
                    else:
                        data = [self.read_legacy_salary_sheet(ws)]
                settings = None
                if 'OT Settings' in wb.sheetnames:
                    settings = {clean_key(r[0]): r[1] for r in wb['OT Settings'].iter_rows(min_row=2, values_only=True) if len(r) > 1 and r[0]}
            finally:
                wb.close()
            data = [d for d in data if d.get('employee_name') or d.get('employee_code') or num_or_zero(d.get('basic', '0')) != 0]
            if sheet != ALL_TAB:                      # the tab name is the company name
                for d in data:
                    if not str(d.get('company', '')).strip(): d['company'] = sheet
            if not data:
                messagebox.showwarning('No Payroll Data', 'No usable employee data was found in that tab.\nTry another tab, or choose "All employees".'); return
            self._loading = True
            try:
                if settings: self.apply_settings(settings)
                if month: self.month.set(month); self.ensure_month_listed()
            finally:
                self._loading = False
            self.set_rows(data)
            self.excel_path = None if is_xls else path; self.sheet_name = sheet
            self.file_label.set(f'{path.name}   ›   {sheet}')
            self.calculate_all(False)
            self.dirty = False; self.set_title()
            self.remember(path, sheet)
            self.status.set(f'Opened {len(data)} employee(s) from "{sheet}". ' + self.status.get())
        except Exception as e:
            self._loading = False
            messagebox.showerror('Open Error', str(e))
        finally:
            self.config(cursor='')

    def sheet_month(self, ws):
        for row in ws.iter_rows(min_row=1, max_row=3, values_only=True):
            if row and clean_key(row[0]) == 'payroll month' and len(row) > 1 and row[1]: return str(row[1]).strip()
        return None

    def active_members_sheet(self, wb):
        return next((n for n in wb.sheetnames if clean_key(n).startswith(ACTIVE_MEMBERS_PREFIX)), None)

    def read_master(self, wb):
        records, master_companies = {}, {}
        if EMPLOYEE_MASTER_SHEET not in wb.sheetnames: return records, master_companies
        ws = wb[EMPLOYEE_MASTER_SHEET]; header = None; hm = {}
        for ri, row in enumerate(ws.iter_rows(values_only=True), 1):
            m = {clean_key(v): i for i, v in enumerate(row) if v is not None}
            if 'empl code' in m and 'employees name' in m: header = ri; hm = m; break
        if header:
            for row in ws.iter_rows(min_row=header+1, values_only=True):
                name = row[hm['employees name']] if hm['employees name'] < len(row) else None
                if not name: continue
                code = row[hm['empl code']] if hm['empl code'] < len(row) else ''
                company = row[hm['working company']] if 'working company' in hm and hm['working company'] < len(row) else ''
                key = clean_person(name)
                d = make_record({'employee_name': str(name).strip(), 'employee_code': str(code or ''), 'company': str(company or '')})
                records[clean_key(code) if code not in (None,'') else key] = d; master_companies[key] = str(company or '')
        return records, master_companies

    def read_active(self, wb, records, master_companies):
        am = self.active_members_sheet(wb)
        if not am: return records
        rows = list(wb[am].iter_rows(values_only=True))
        if not rows: return records
        hm = {clean_key(v): i for i, v in enumerate(rows[0]) if v is not None}
        for row in rows[1:]:
            ni = hm.get('name'); ci = hm.get('employee code')
            name = row[ni] if ni is not None and ni < len(row) else None
            code = row[ci] if ci is not None and ci < len(row) else None
            if not name: continue
            key = clean_person(name); codekey = clean_key(code) if code not in (None,'') else ''
            d = records.get(codekey) if codekey else None
            if d is None: d = next((v for v in records.values() if clean_person(v.get('employee_name','')) == key), None)
            if d is None:
                d = make_record({'employee_name': str(name).strip(), 'employee_code': str(code or ''), 'company': master_companies.get(key, '')})
            else:
                d['employee_name'] = str(name).strip()
                if code not in (None,''): d['employee_code'] = re.sub(r'^(\d+)\.0+$', r'\1', str(code))
            records[codekey or key] = d
        return records

    def read_all_reference_employees(self, wb):
        records, master_companies = self.read_master(wb)
        records = self.read_active(wb, records, master_companies)
        for sn in SALARY_SHEETS:
            if sn not in wb.sheetnames: continue
            salary = self.read_legacy_salary_sheet(wb[sn]); skey = clean_person(salary.get('employee_name',''))
            target = next((d for d in records.values() if clean_person(d.get('employee_name','')) == skey), None)
            if target is None and salary.get('employee_name'):
                target = make_record(salary)
                records[clean_key(target.get('employee_code')) or skey] = target
            elif target is not None:
                target.update({k: v for k, v in make_record(salary).items() if k not in ('employee_name','employee_code','company') and v not in ('','0')})
        out = []; seen = set()
        for d in records.values():
            key = clean_key(d.get('employee_code')) or clean_person(d.get('employee_name',''))
            if not key or key in seen: continue
            seen.add(key); out.append(d)
        return out

    def looks_like_export(self, ws):
        for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
            vals = [str(c or '').strip().lower() for c in row]
            if 'employee name' in vals or 'net salary' in vals: return True
        return False

    def read_tabular_sheet(self, ws):
        rows = list(ws.iter_rows(values_only=True)); header_idx = None; mapping = {}
        for i, row in enumerate(rows[:30]):
            for j, v in enumerate(row):
                k = clean_key(v)
                if k in HEADER_ALIASES: mapping[j] = HEADER_ALIASES[k]
            if {'employee_name', 'employee_code'} & set(mapping.values()): header_idx = i; break   # mapping is {column: field}
        if header_idx is None: return []
        out = []
        for row in rows[header_idx+1:]:
            d = {}
            for j, key in mapping.items():
                if j < len(row) and row[j] is not None: d[key] = str(row[j])
            rec = make_record(d)
            if any(rec[k] not in ('', '0', '26') for k in ['employee_name', 'employee_code', 'basic']): out.append(rec)
        return out

    def read_legacy_salary_sheet(self, ws):
        """Read the older one-employee salary-slip layout (labels in column B, values from column C)."""
        d = dict(DEFAULTS); rows = list(ws.iter_rows(values_only=True))
        for row in rows[:8]:
            val = row[2] if len(row) > 2 else None
            if isinstance(val, str) and val.strip() and val.strip().upper() not in ('PER MONTH', 'PER DAY', 'NO.DAYS'):
                d['employee_name'] = val.strip(); break
        d.update({'working_days': '26', 'paid_holidays': '0', 'ot_hours': '0'})
        basic_set = False
        for row_index, row in enumerate(rows[:25], start=1):
            label = str(row[1] or '').strip().upper() if len(row) > 1 else ''
            vals = list(row[2:]); first = next((x for x in vals if x is not None), None)
            monthly = row[4] if len(row) > 4 and row[4] is not None else first
            isnum = isinstance(monthly, (int, float))
            if label in ('W.DAYS', 'W DAYS') and first is not None: d['working_days'] = str(first)
            elif label == 'PH' and first is not None: d['paid_holidays'] = str(first)
            elif label in ('O.T HRS', 'OT HRS') and first is not None: d['ot_hours'] = str(first)
            elif label == 'BASIC': d['basic'] = str(monthly or 0); basic_set = True
            elif not label and not basic_set and isnum and 8 <= row_index <= 12:
                d['basic'] = str(monthly); basic_set = True       # unlabeled first pay row = basic
            elif label in ('WASING ALLOW', 'WASHING ALLOW', 'WASHING ALLOWANCE'): d['allowances'] = str(monthly or 0)
            elif label in ('HRA', 'HAMALI CHARGES', 'PRODUCTION INCENTIVE'): d['other_earnings'] = str(num_or_zero(d['other_earnings']) + num_or_zero(monthly or 0))
            elif label == 'OT' and isnum:
                hrs = num_or_zero(d['ot_hours'])
                d['ot_rate'] = str((num(monthly) / hrs).quantize(Decimal('0.01'))) if hrs else '0'
            elif label == 'ESI' and isinstance(monthly, str) and re.search(r'EX[AE]MPT', monthly.upper()): d['esi_exempt'] = 'Yes'   # sheet spells it "EXAMPTED"
            elif label == 'ADVANCE' and isnum: d['festival_advance'] = str(monthly)
            elif label == 'GNERAL' and isnum: d['general_recovery'] = str(monthly)
            elif label == 'F.ADVANCE' and isnum: d['festival_recovery'] = str(monthly)
        return make_record(d)

    def convert_legacy_xls(self, src, dst):
        for exe in ('libreoffice', 'soffice'):
            try:
                subprocess.run([exe, '--headless', '--convert-to', 'xlsx', '--outdir', str(dst.parent), str(src)],
                               stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60)
                generated = dst.parent / (Path(src).stem + '.xlsx')
                if generated.exists(): generated.replace(dst); return True
            except Exception: pass
        return False

    # ------------------------------------------------------------------ grid: rows
    def new_iid(self):
        self.next_id += 1; return f'r{self.next_id}'

    def set_rows(self, records):
        """Replace the grid with `records` (list of dicts); 5 blank rows if empty."""
        self.cancel_editor()
        self._loading = True
        try:
            self.tree.delete(*self.tree.get_children())
            self.data, self.flags = {}, {}
            for rec in (records or [make_record() for _ in range(5)]):
                self.insert_record(make_record(rec))
        finally:
            self._loading = False
        self.tree.xview_moveto(0); self.tree.yview_moveto(0)

    def insert_record(self, rec):
        iid = self.new_iid(); self.data[iid] = rec; self.flags[iid] = None
        idx = len(self.tree.get_children())
        self.tree.insert('', 'end', iid=iid, values=self.display_values(iid, idx), tags=('odd' if idx % 2 == 0 else 'even',))
        return iid

    def display_values(self, iid, idx):
        d = self.data[iid]
        return [idx + 1] + ['[ Advance ]' if k == 'advance_btn' else d[k] for k in GRID_KEYS]

    def refresh_row(self, iid):
        if iid not in self.data: return
        idx = self.tree.index(iid)
        tag = self.flags.get(iid) or ('odd' if idx % 2 == 0 else 'even')
        self.tree.item(iid, values=self.display_values(iid, idx), tags=(tag,))

    def refresh_all(self):
        for iid in self.tree.get_children(): self.refresh_row(iid)

    def add_employee(self):
        iid = self.insert_record(make_record()); self.mark_dirty()
        self.tree.selection_set(iid); self.tree.focus(iid); self.tree.see(iid)
        self.recalc_row(iid)
        self.update_idletasks(); self.start_edit(iid, 'employee_name')

    def selected_iid(self):
        sel = self.tree.selection()
        return sel[0] if sel else None

    def delete_selected(self):
        iid = self.selected_iid()
        if iid is None:
            messagebox.showinfo('Select an employee', 'Click a row first, then press DELETE SELECTED.'); return
        d = self.data[iid]; label = d['employee_name'] or d['employee_code'] or 'this row'
        if not messagebox.askyesno('Delete employee', f'Delete {label} from the grid?\n(The saved Excel file is not changed until you SAVE.)'): return
        nxt = self.tree.next(iid) or self.tree.prev(iid)
        self.tree.delete(iid); self.data.pop(iid, None); self.flags.pop(iid, None)
        if nxt: self.tree.selection_set(nxt)
        self.refresh_all(); self.mark_dirty(); self.calculate_all(False); self.status.set(f'Deleted {label}.')

    def clear_all(self):
        if not messagebox.askyesno('Clear all', 'Remove every employee from the grid?\n(The saved Excel file is not changed until you SAVE.)'): return
        self.set_rows([]); self.error_rows = []; self.mark_dirty(); self.update_summary(); self.status.set('All rows cleared.')

    def on_select(self, event=None):
        iid = self.selected_iid()
        if iid and iid in self.data:
            d = self.data[iid]
            self.banner.set(f'Selected: {d["employee_name"] or "(no name)"}  (code {d["employee_code"] or "-"})   |   click [ Advance ] or the Advance button for advance details')

    # ------------------------------------------------------------------ grid: cell editing
    def on_click(self, ev):
        region = self.tree.identify_region(ev.x, ev.y)
        if region != 'cell':
            self.commit_editor(); return
        iid = self.tree.identify_row(ev.y)
        col = self.tree.identify_column(ev.x)
        if not iid or not col: return
        key = TREE_COLS[int(col[1:]) - 1]
        self.commit_editor()
        self.tree.selection_set(iid); self.tree.focus(iid)
        if key == 'advance_btn':
            self.after_idle(lambda: self.open_advance(iid)); return 'break'
        if key in EDIT_KEYS:
            self.after(10, lambda: self.start_edit(iid, key, auto_open=True)); return 'break'
        # calculated columns / row number: just select the row

    def on_tree_return(self, ev):
        iid = self.selected_iid()
        if iid: self.start_edit(iid, EDIT_KEYS[0])
        return 'break'

    def start_edit(self, iid, key, auto_open=False):
        self.commit_editor()
        if iid not in self.data or key not in EDIT_KEYS: return
        self.reveal(iid, key)
        self.tree.update_idletasks()
        box = self.tree.bbox(iid, key)
        if not box: return
        self.tree.selection_set(iid); self.tree.focus(iid)
        typ = GRID_TYPE[key]; value = self.data[iid][key]
        if typ in ('choice', 'yesno'):
            e = ttk.Combobox(self.tree, state='readonly', values=OT_METHODS if typ == 'choice' else ['No', 'Yes'], font=(self.fam, 10))
            e.set(value)
            e.bind('<<ComboboxSelected>>', lambda ev: (self.commit_editor(), self.move(iid, key, 'next')))
            e.bind('<space>', lambda ev: self._post(e))
            e.bind('<Return>', lambda ev: self._finish(iid, key, 'next'))
        else:
            e = tk.Entry(self.tree, font=(self.fam, 10), relief='solid', borderwidth=1, bg='#fffbe6', highlightthickness=2,
                         highlightcolor=ACCENT, highlightbackground=ACCENT, justify='left' if typ == 'text' else 'right')
            e.insert(0, value)
            e.bind('<Return>', lambda ev: self._finish(iid, key, 'next'))
            e.bind('<FocusOut>', lambda ev, e=e: self.commit_editor() if self.editor is e else None)
        e.bind('<Tab>', lambda ev: self._finish(iid, key, 'next'))
        e.bind('<Shift-Tab>', lambda ev: self._finish(iid, key, 'prev'))
        e.bind('<ISO_Left_Tab>', lambda ev: self._finish(iid, key, 'prev'))
        e.bind('<Down>', lambda ev: self._finish(iid, key, 'down') if typ not in ('choice', 'yesno') else None)
        e.bind('<Up>', lambda ev: self._finish(iid, key, 'up') if typ not in ('choice', 'yesno') else None)
        e.bind('<Escape>', lambda ev: (self.cancel_editor(), self.tree.focus_set(), 'break')[-1])
        self.editor, self.cur = e, (iid, key)
        e.place(x=box[0], y=box[1], width=box[2], height=box[3])
        e.focus_set()
        if isinstance(e, tk.Entry): e.selection_range(0, 'end'); e.icursor('end')
        elif auto_open: self.after(60, lambda: self._post(e))

    def _post(self, combo):
        try: combo.tk.call('ttk::combobox::Post', str(combo))
        except tk.TclError: pass
        return 'break'

    def _finish(self, iid, key, direction):
        self.commit_editor(); self.move(iid, key, direction); return 'break'

    def move(self, iid, key, direction):
        """Enter/Tab = next editable cell (then the next row); Shift+Tab = previous; Up/Down = same column."""
        rows = list(self.tree.get_children())
        if iid not in rows: return
        r, c = rows.index(iid), EDIT_KEYS.index(key)
        if direction == 'next':
            c += 1
            if c >= len(EDIT_KEYS): c = 0; r += 1
        elif direction == 'prev':
            c -= 1
            if c < 0: c = len(EDIT_KEYS) - 1; r -= 1
        elif direction == 'down': r += 1
        elif direction == 'up': r -= 1
        if r < 0: return
        if r >= len(rows):
            self.status.set('Last cell reached. Use ADD EMPLOYEE for a new row.'); self.tree.focus_set(); return
        self.start_edit(rows[r], EDIT_KEYS[c])

    def reveal(self, iid, key):
        """Scroll so that the cell is fully visible."""
        self.tree.see(iid)
        cols = TREE_COLS; widths = [int(self.tree.column(k, 'width')) for k in cols]
        total = sum(widths) or 1
        x0 = sum(widths[:cols.index(key)]); x1 = x0 + widths[cols.index(key)]
        self.tree.update_idletasks()
        view = max(self.tree.winfo_width(), 1); left = self.tree.xview()[0] * total
        if x0 < left: self.tree.xview_moveto(x0 / total)
        elif x1 > left + view: self.tree.xview_moveto(min(1.0, (x1 - view + 4) / total))

    def _yscroll(self, *a):
        self.vs.set(*a); self._reposition_editor()

    def _xscroll(self, *a):
        self.hs.set(*a); self._reposition_editor()

    def _reposition_editor(self):
        if not self.editor or not self.cur: return
        try:
            box = self.tree.bbox(*self.cur)
            if box: self.editor.place(x=box[0], y=box[1], width=box[2], height=box[3])
            else: self.editor.place_forget()
        except tk.TclError: pass

    def cancel_editor(self):
        e, self.editor, self.cur = self.editor, None, None
        if e is not None:
            try: e.destroy()
            except tk.TclError: pass

    def commit_editor(self):
        e, cur = self.editor, self.cur
        if e is None or cur is None: return
        try: text = e.get().strip()
        except tk.TclError: text = None
        self.editor, self.cur = None, None
        try: e.destroy()
        except tk.TclError: pass
        if text is not None: self.set_value(cur[0], cur[1], text)

    def set_value(self, iid, key, text):
        if iid not in self.data: return
        typ = GRID_TYPE[key]
        if key in WHOLE_INPUTS and text:
            whole = to_whole(text)
            if whole != text and whole.lstrip('-').isdigit():
                self.status.set(f'{LABEL[key]} must be a whole number - {text} was rounded to {whole}.'); text = whole
        if typ == 'num' and text == '': text = '0'
        if self.data[iid][key] == text: return
        self.data[iid][key] = text; self.mark_dirty()
        self.recalc_row(iid)

    # ------------------------------------------------------------------ advance pop-up
    def open_selected_advance(self):
        iid = self.selected_iid()
        if iid is None:
            messagebox.showinfo('Select an employee', 'Click an employee row first, then press ADVANCE.'); return
        self.open_advance(iid)

    def open_advance(self, iid):
        if iid not in self.data: return
        AdvanceDialog(self, iid)

    def advance_saved(self, iid):
        self.mark_dirty(); self.recalc_row(iid)
        d = self.data[iid]
        self.status.set(f'Advance saved for {d["employee_name"] or d["employee_code"]}: deduction {money(d["advance_deduction"])}, balance {money(d["advance_balance"])}.')

    # ------------------------------------------------------------------ calculation
    def row_data(self, iid): return self.data[iid]

    @staticmethod
    def _nonblank_num(s):
        try: return num(s) != 0
        except ValueError: return True

    def is_active(self, iid):
        v = self.data[iid]
        if v['employee_name'].strip() or v['employee_code'].strip(): return True
        return any(self._nonblank_num(v[k]) for k in MONEY_NONZERO_KEYS)

    def active_ids(self): return [i for i in self.tree.get_children() if self.is_active(i)]
    def active_rows(self): return [self.data[i] for i in self.active_ids()]

    def calculate_row(self, iid, settings=None):
        d = self.data[iid]
        try:
            calc, warns = compute(d, settings if settings is not None else self.settings())
        except ValueError:
            self.flags[iid] = 'bad'; self.refresh_row(iid); raise
        d.update(calc)
        self.flags[iid] = 'bad' if any(w[0] == 'error' for w in warns) else ('warn' if warns else None)
        self.refresh_row(iid)
        return calc, warns

    def recalc_row(self, iid):
        if self._loading or iid not in self.data: return
        try: _, warns = self.calculate_row(iid)
        except ValueError as e: self.status.set(f'Row {self.tree.index(iid) + 1}: {e}')
        else:
            self.status.set(' | '.join(w[1] for w in warns[:2]) if warns else 'Row updated.')
        self.update_summary()

    def update_summary(self):
        g = ded = net = D('0'); n = 0
        for iid in self.active_ids():
            d = self.data[iid]
            gross, net_v = num_or_zero(d['gross']), num_or_zero(d['net_salary'])
            if gross == 0 and net_v == 0: continue
            g += gross; ded += num_or_zero(d['total_deductions']); net += net_v; n += 1
        self.sum_vars['count'].set(str(n)); self.sum_vars['gross'].set(money(g))
        self.sum_vars['ded'].set(money(ded)); self.sum_vars['net'].set(money(net))

    def calculate_all(self, show=True):
        self.commit_editor()
        try:
            s = self.settings()
        except ValueError as e:
            self.status.set(str(e))
            if show: messagebox.showerror('Settings Error', str(e))
            return False
        totals = [D('0')] * 3; count = 0; msgs = []; self.error_rows = []; bad = []
        seen = {}
        for iid in self.active_ids():
            d = self.data[iid]; who = d['employee_name'] or d['employee_code'] or f'row {self.tree.index(iid) + 1}'
            try:
                calc, warns = self.calculate_row(iid, s)
            except ValueError as e:
                bad.append(f'{who}: {e}'); continue
            totals = [totals[0] + num(calc['gross']), totals[1] + num(calc['total_deductions']), totals[2] + num(calc['net_salary'])]; count += 1
            for lvl, msg, _ in warns:
                msgs.append(f'{who}: {msg}')
                if lvl == 'error': self.error_rows.append(who)
            code = d['employee_code'].strip().lower()
            if code:
                if code in seen: msgs.append(f'{who}: duplicate employee code {d["employee_code"]} (also {seen[code]}).')
                else: seen[code] = who
        self.update_summary()
        self.status.set(f'Calculated {count} employee(s) | Gross {money(totals[0])} | Deductions {money(totals[1])} | Net {money(totals[2])}'
                        + (f' | {len(msgs)} warning(s) (rows shown red / yellow)' if msgs else ''))
        if bad:
            if show: messagebox.showerror('Calculation Error', '\n'.join(bad[:10]))
            self.status.set(f'Input error - {bad[0]}'); return False
        if show and msgs:
            more = f'\n...and {len(msgs) - 15} more.' if len(msgs) > 15 else ''
            messagebox.showwarning('Please review', '\n'.join(msgs[:15]) + more)
        self.warnings = msgs
        return True

    def confirm_errors(self):
        """Ask before saving/printing when a row has an error-level problem (e.g. negative net salary)."""
        if not self.error_rows: return True
        names = ', '.join(dict.fromkeys(self.error_rows))
        return messagebox.askyesno('Problems found', f'These employees have errors (e.g. negative net salary):\n{names}\n\nContinue anyway?')

    # ------------------------------------------------------------------ Excel save
    def save_to_excel(self):
        if not OPENPYXL: messagebox.showerror('Missing Package', 'Install openpyxl: pip install openpyxl'); return False
        if not self.calculate_all(): return False
        active = self.active_rows()
        if not active: messagebox.showwarning('No Employees', 'Enter at least one employee.'); return False
        if not self.confirm_errors(): return False
        if self.excel_path is None:
            path = filedialog.asksaveasfilename(defaultextension='.xlsx', initialfile=f'Payroll_{self.month.get().replace(" ","_")}.xlsx', filetypes=[('Excel Workbook','*.xlsx')])
            if not path: return False
            self.excel_path = Path(path)
        self.config(cursor='watch'); self.update_idletasks()
        try:
            self.write_excel(self.excel_path, active)
        except PermissionError:
            messagebox.showerror('Cannot save', f'{self.excel_path.name} is open in another program (Excel?). Close it and try again.'); return False
        except Exception as e:
            messagebox.showerror('Save Error', str(e)); return False
        finally:
            self.config(cursor='')
        self.dirty = False; self.set_title()
        self.file_label.set(f'{self.excel_path.name}   ›   Payroll Data')
        self.remember(self.excel_path, 'Payroll Data')
        self.status.set(f'Saved {len(active)} employee(s) to {self.excel_path}')
        messagebox.showinfo('Saved to Excel', f'{len(active)} employee record(s) saved to:\n{self.excel_path}\n\nTabs updated: "Payroll Data" and the month tab.\nA backup of the previous file is in the "backups" folder.')
        return True

    def backup_workbook(self, path):
        path = Path(path)
        if not path.exists(): return
        BACKUP_DIR.mkdir(exist_ok=True)
        shutil.copy2(path, BACKUP_DIR / f'{path.stem}_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}')
        old = sorted(BACKUP_DIR.glob(f'{path.stem}_*{path.suffix}'))
        for f in old[:-MAX_BACKUPS]:
            try: f.unlink()
            except OSError: pass

    def typed_row(self, d):
        out = []
        for _, k, _, typ in COLUMNS:
            if typ in ('num', 'calc'):
                try:
                    n = num(d[k]); out.append(int(n) if k in WHOLE_FIELDS else float(n))
                except ValueError: out.append(d[k])
            else:
                value = d[k]
                # Keep user-entered text from becoming an Excel formula on export.
                if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
                    value = "'" + value
                out.append(value)
        return out

    def write_sheet(self, wb, title, active):
        if title in wb.sheetnames: wb.remove(wb[title])
        ws = wb.create_sheet(title)
        ws.append(['Payroll Month', self.month.get()]); ws.append(['Company Working Days', self.company_working_days.get()]); ws.append([])
        ws.append([c[0] for c in COLUMNS])
        thin = Side(style='thin', color='B7B7B7')
        for cell in ws[4]:
            cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='4472C4')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = Border(bottom=thin)
        for d in active: ws.append(self.typed_row(d))
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(max(max(len(str(c.value or '')) for c in col[3:]) + 2, 12), 28)
        ws.freeze_panes = 'A5'; ws.auto_filter.ref = f'A4:{ws.cell(ws.max_row, ws.max_column).coordinate}'
        return ws

    def write_excel(self, path, active):
        path = Path(path)
        self.backup_workbook(path)
        if path.exists(): wb = load_workbook(path)
        elif REFERENCE_XLSX.exists(): shutil.copy2(REFERENCE_XLSX, path); wb = load_workbook(path)
        else: wb = Workbook(); wb.remove(wb.active)
        self.write_sheet(wb, 'Payroll Data', active)
        # Keep one snapshot sheet per month so earlier months are not lost when the next month is saved.
        snap = re.sub(r'[\[\]:*?/\\]', '-', f'Payroll {self.month.get().strip()}')[:31] or 'Payroll Snapshot'
        if snap != 'Payroll Data': self.write_sheet(wb, snap, active)
        if 'OT Settings' in wb.sheetnames: wb.remove(wb['OT Settings'])
        ws2 = wb.create_sheet('OT Settings'); ws2.append(['Setting', 'Value'])
        for row in self.settings_rows(): ws2.append(list(row))
        for c in ws2[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='4472C4')
        ws2.column_dimensions['A'].width = 38; ws2.column_dimensions['B'].width = 44
        self.sync_master(wb, active)
        temp = path.with_name(path.stem + '.tmp.xlsx'); wb.save(temp); temp.replace(path)

    def sync_master(self, wb, active):
        """Append employees that are not yet on the employee master sheet."""
        if EMPLOYEE_MASTER_SHEET not in wb.sheetnames: return
        master = wb[EMPLOYEE_MASTER_SHEET]; header_row = None; hmap = {}
        for ri, row in enumerate(master.iter_rows(values_only=True), 1):
            m = {clean_key(v): i+1 for i, v in enumerate(row) if v is not None}
            if 'empl code' in m and 'employees name' in m: header_row = ri; hmap = m; break
        if not header_row: return
        code_col, name_col = hmap['empl code'], hmap['employees name']; company_col = hmap.get('working company')
        names = set(); codes = set()
        for row in master.iter_rows(min_row=header_row+1, values_only=True):
            if len(row) >= name_col and row[name_col-1]: names.add(clean_person(row[name_col-1]))
            if len(row) >= code_col and row[code_col-1] not in (None, ''): codes.add(clean_key(row[code_col-1]))
        for d in active:
            name = d['employee_name'].strip(); code = d['employee_code'].strip()
            if not name or clean_person(name) in names or (code and clean_key(code) in codes): continue
            vals = [None] * master.max_column; vals[code_col-1] = code or None; vals[name_col-1] = name
            if company_col and company_col <= len(vals): vals[company_col-1] = d.get('company') or None
            master.append(vals); names.add(clean_person(name))
            if code: codes.add(clean_key(code))

    # ------------------------------------------------------------------ exit
    def on_exit(self):
        self.commit_editor()
        if self.dirty and self.active_ids():
            ans = messagebox.askyesnocancel('Save before exit?', 'You have unsaved changes.\n\nYes = save to Excel and exit\nNo = exit without saving\nCancel = stay')
            if ans is None: return
            if ans and not self.save_to_excel(): return
        self.destroy()

    # ------------------------------------------------------------------ PDF
    def make_pdf(self, path):
        font, bold, uni = pdf_fonts(); sym = '₹' if uni else 'Rs.'
        rows = self.active_rows()
        companies = ', '.join(dict.fromkeys(d['company'].strip() for d in rows if d['company'].strip()))
        paid = [d for d in rows if num_or_zero(d['gross']) != 0 or num_or_zero(d['net_salary']) != 0]
        skipped = len(rows) - len(paid)
        page = landscape(A4); L = 18; usable = page[0] - 2 * L
        doc = SimpleDocTemplate(path, pagesize=page, leftMargin=L, rightMargin=L, topMargin=22, bottomMargin=26,
                                title='Payroll Statement', author='Payroll Master')
        styles = getSampleStyleSheet()
        title = ParagraphStyle('T', parent=styles['Title'], alignment=TA_CENTER, fontSize=16, fontName=bold)
        norm = ParagraphStyle('N', parent=styles['Normal'], fontName=font, fontSize=8)
        cell = ParagraphStyle('C', parent=styles['Normal'], fontName=font, fontSize=6.3, leading=7.4)
        story = [Paragraph('MULTIPLE EMPLOYEE PAYROLL STATEMENT', title), Spacer(1, 4),
                 Paragraph((f'<b>Company:</b> {companies.replace("&", "&amp;")} &nbsp;&nbsp; ' if companies else '') +
                           f'<b>Payroll Month:</b> {self.month.get()} &nbsp;&nbsp; <b>Company Working Days:</b> {self.company_working_days.get()}'
                           f' &nbsp;&nbsp; <b>Employees paid:</b> {len(paid)} &nbsp;&nbsp; <b>Generated:</b> {datetime.now():%d-%m-%Y %H:%M}', norm), Spacer(1, 8)]
        text_keys = {'employee_name': 11, 'employee_code': 5, 'company': 10, 'ot_method': 12}
        plain_keys = {'working_days', 'paid_holidays', 'paid_days', 'multiplier', 'hours_per_day', 'ot_hours', 'esi_exempt'}
        sections = [('Employee / Earnings', ['employee_name','employee_code','working_days','paid_holidays','paid_days','basic','allowances',
                                              'other_earnings','ot_method','multiplier','hours_per_day','ot_hours','ot_rate','ot_amount','gross']),
                    ('Advances / Deductions / Net Pay', ['employee_name','employee_code','festival_advance','festival_recovery','festival_balance',
                                                         'general_advance','general_recovery','previous_balance','general_balance','pf','epf','eps',
                                                         'esi_exempt','esi','other_deductions','total_deductions','net_salary'])]
        total_keys = {'basic','allowances','other_earnings','ot_amount','gross','festival_advance','festival_recovery','festival_balance','general_advance',
                      'general_recovery','general_balance','pf','epf','eps','esi','other_deductions','total_deductions','net_salary'}
        for name, keys in sections:
            weights = [text_keys.get(k, 4.2 if k in plain_keys else 6.4) for k in keys]
            widths = [w / sum(weights) * usable for w in weights]
            head = [Paragraph(f'<b>{LABEL[k]}</b>', ParagraphStyle('H', parent=cell, alignment=TA_CENTER)) for k in keys]
            data = [[Paragraph(f'<b>{name}</b>', ParagraphStyle('S', parent=cell, textColor=colors.white, fontSize=7.5))] + [''] * (len(keys) - 1), head]
            for d in paid:
                line = []
                for k in keys:
                    if k in text_keys: line.append(Paragraph(str(d[k]).replace('&', '&amp;'), cell))
                    elif k == 'esi_exempt': line.append(d[k])
                    elif k in plain_keys: line.append(plain(d[k]))
                    else: line.append(money(d[k], sym))
                data.append(line)
            tot = []
            for k in keys:
                if k == keys[0]: tot.append('TOTAL')
                elif k in total_keys: tot.append(money(sum((num_or_zero(d[k]) for d in paid), D('0')), sym))
                else: tot.append('')
            data.append(tot)
            first_num = next(i for i, k in enumerate(keys) if k not in text_keys)
            t = Table(data, colWidths=widths, repeatRows=2)
            t.setStyle(TableStyle([('SPAN',(0,0),(-1,0)), ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#4472C4')), ('BACKGROUND',(0,1),(-1,1),colors.HexColor('#DDE5F3')),
                                   ('GRID',(0,0),(-1,-1),.35,colors.grey), ('FONTNAME',(0,0),(-1,-1),font), ('FONTSIZE',(0,0),(-1,-1),6.3),
                                   ('FONTNAME',(0,-1),(-1,-1),bold), ('BACKGROUND',(0,-1),(-1,-1),colors.HexColor('#EEEEEE')),
                                   ('ALIGN',(first_num,2),(-1,-1),'RIGHT'), ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
                                   ('ROWBACKGROUNDS',(0,2),(-1,-2),[colors.white, colors.HexColor('#F7F9FC')]),
                                   ('LEFTPADDING',(0,0),(-1,-1),2.5), ('RIGHTPADDING',(0,0),(-1,-1),2.5), ('TOPPADDING',(0,0),(-1,-1),2.5), ('BOTTOMPADDING',(0,0),(-1,-1),2.5)]))
            story += [t, Spacer(1, 10)]
        note = ('<b>Note:</b> Advance "Given" amounts are outstanding employee loans and do not reduce current Net Salary. Only the recovery amounts '
                '(General and Festival) that are not deferred reduce Net Salary. A recovery marked "Defer" is not deducted this month. '
                'EPF/EPS parts show how the PF amount is split and are not deducted again.')
        if skipped: note += f' {skipped} employee(s) with no salary entered were left out of this statement.'
        story.append(Paragraph(note, norm))
        def footer(canvas, doc_):
            canvas.saveState(); canvas.setFont(font, 7); canvas.drawRightString(page[0] - L, 12, f'Page {doc_.page}'); canvas.restoreState()
        doc.build(story, onFirstPage=footer, onLaterPages=footer)

    def _pdf_precheck(self):
        if not REPORTLAB: messagebox.showerror('Missing Package', 'Install reportlab:  pip install reportlab'); return False
        if not self.active_rows(): messagebox.showwarning('No Employees', 'Enter at least one employee.'); return False
        if not self.calculate_all(): return False
        return self.confirm_errors()

    def pdf_all(self):
        if not self._pdf_precheck(): return
        path = filedialog.asksaveasfilename(defaultextension='.pdf', initialfile=f'Payroll_{self.month.get().replace(" ","_")}.pdf', filetypes=[('PDF','*.pdf')])
        if path:
            try: self.make_pdf(path); self.status.set('Payroll PDF generated successfully.'); messagebox.showinfo('PDF Generated', path)
            except Exception as e: messagebox.showerror('PDF Error', str(e))

    def print_all(self):
        if not self._pdf_precheck(): return
        path = filedialog.asksaveasfilename(defaultextension='.pdf', initialfile=f'Payroll_{self.month.get().replace(" ","_")}_print.pdf', filetypes=[('PDF','*.pdf')])
        if not path: return
        try:
            self.make_pdf(path)
            if sys.platform.startswith('win'): os.startfile(path, 'print')
            elif sys.platform == 'darwin': subprocess.Popen(['open', path]); messagebox.showinfo('Print', 'PDF opened. Use Command+P to print.')
            else: subprocess.Popen(['xdg-open', path]); messagebox.showinfo('Print', 'PDF opened. Use Ctrl+P to print.')
        except Exception as e: messagebox.showerror('Print Error', str(e))


if __name__ == '__main__':
    App().mainloop()
